from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_project_plan import (
    COLOR_NAVY,
    COLOR_TEXT_GRAY,
    FONT_BODY,
    FONT_HEAD,
    add_bullet,
    add_heading,
    add_hyperlink,
    add_paragraph,
    add_table,
    set_run_font,
)
from build_project_plan_v02 import add_metadata_table


SOURCE_PATH = Path("deliverables/个人食品决策助手第一阶段项目落实计划书_v0.2.docx")
OUTPUT_PATH = Path("deliverables/个人食品决策助手第一阶段项目落实计划书_修订3.docx")


def clear_paragraph(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def reset_footer_page_number(section):
    paragraph = section.footer.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, FONT_BODY, 8.5, False, COLOR_TEXT_GRAY)

    # A two-character cached result prevents LibreOffice from allocating only
    # one digit of width when the dynamic PAGE field reaches page 11 or later.
    result_run = paragraph.add_run("88")
    set_run_font(result_run, FONT_BODY, 8.5, False, COLOR_TEXT_GRAY)
    paragraph._p.remove(result_run._r)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field.append(result_run._r)
    paragraph._p.append(field)

    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, FONT_BODY, 8.5, False, COLOR_TEXT_GRAY)


def set_paragraph_text(paragraph, text):
    clear_paragraph(paragraph)
    run = paragraph.add_run(text)
    style_name = paragraph.style.name if paragraph.style is not None else "Normal"
    if style_name == "Title":
        set_run_font(run, FONT_HEAD, 25, True, "000000")
    elif style_name == "Document Subtitle":
        set_run_font(run, FONT_BODY, 12, False, COLOR_TEXT_GRAY)
    elif style_name == "Heading 1":
        set_run_font(run, FONT_HEAD, 15, True, "000000")
    elif style_name == "Heading 2":
        set_run_font(run, FONT_HEAD, 12.2, True, "000000")
    else:
        set_run_font(run, FONT_BODY, 10.8, False, "000000")


def find_paragraph(doc, text):
    for paragraph in doc.paragraphs:
        if paragraph.text == text:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def replace_paragraph(doc, old, new):
    paragraph = find_paragraph(doc, old)
    set_paragraph_text(paragraph, new)
    return paragraph


def find_table(doc, headers):
    for table in doc.tables:
        if not table.rows:
            continue
        values = [cell.text.strip() for cell in table.rows[0].cells]
        if values == list(headers):
            return table
    raise ValueError(f"Table not found: {headers}")


def remove_element(element):
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def replace_table(doc, old_headers, new_headers, rows, widths, center_cols, font_size):
    old_table = find_table(doc, old_headers)
    new_table = add_table(
        doc,
        new_headers,
        rows,
        widths=widths,
        center_cols=center_cols,
        font_size=font_size,
    )
    trailing = new_table._tbl.getnext()
    old_table._tbl.addprevious(new_table._tbl)
    remove_element(old_table._tbl)
    if trailing is not None and trailing.tag == qn("w:p"):
        remove_element(trailing)
    return new_table


def replace_metadata_table(doc, rows):
    old_table = find_table(doc, ["项目概况", "内容"])
    before_tail = doc._element.body[-2]
    add_metadata_table(doc, rows)
    new_table = None
    cursor = before_tail.getnext()
    added = []
    while cursor is not None and cursor.tag != qn("w:sectPr"):
        next_cursor = cursor.getnext()
        added.append(cursor)
        if cursor.tag == qn("w:tbl") and new_table is None:
            new_table = cursor
        cursor = next_cursor
    if new_table is None:
        raise RuntimeError("Metadata table was not created")
    old_table._tbl.addprevious(new_table)
    remove_element(old_table._tbl)
    for element in added:
        if element is not new_table:
            remove_element(element)


def insert_block_before(doc, anchor_paragraph, builder):
    body = doc._element.body
    section_properties = body[-1]
    previous = section_properties.getprevious()
    builder()
    cursor = previous.getnext()
    added = []
    while cursor is not None and cursor is not section_properties:
        next_cursor = cursor.getnext()
        added.append(cursor)
        cursor = next_cursor
    for element in added:
        anchor_paragraph._p.addprevious(element)


def build_document():
    doc = Document(SOURCE_PATH)
    reset_footer_page_number(doc.sections[0])

    # Front matter and revision naming.
    set_paragraph_text(doc.paragraphs[0], "个人食品决策助手第一阶段项目落实计划书")
    set_paragraph_text(doc.paragraphs[1], "计划书修订 3  产品目标 v0.1  原生 iOS 个人自用项目")
    replace_metadata_table(
        doc,
        [
            ("计划书修订", "修订 3"),
            ("修订日期", "2026 年 9 月 30 日"),
            ("产品目标", "v0.1"),
            ("目标设备", "iPhone 16 Pro Max、iOS 26.6.2、Apple Watch"),
            ("用户目标", "健身小白阶段的减脂增肌，同时处理 LDL 约束、异维 A 酸疗程和清真饮食"),
            ("第一阶段周期", "6 周开发，加 2 周个人试用"),
        ],
    )

    replace_paragraph(
        doc,
        "第一阶段先完成每天都会使用的记录闭环，再完成超市购买判断。第 2 周交付家常菜称重和手工外卖估算，第 3 周交付目标与今日页，从这时开始积累真实数据。第 4 至第 6 周完成标签扫描、约束判断、HealthKit 和导出。开发结束后单独安排 2 周试用，不再把试用压进第 6 周。",
        "第一阶段仍按 6 周开发加 2 周试用推进。修订 3 不增加新的产品模块，重点把开工前发现的规则缺口补齐：维生素 A 补剂采用更严格的红灯条件，营养分给出可编码公式，油脂和酱料改用适合其食用方式的评分依据，并把 OCR 和 HealthKit 的不确定性提前到第 1 周验证。",
    )
    replace_paragraph(
        doc,
        "购买判断分成食品、补剂和清真三个维度。系统先查找明确的违规证据，再检查信息完整度；补剂使用独立结论，清真状态显示为独立标记。未完成人工核对的扫描结果只能显示临时结论，不得显示绿灯。",
        "购买判断继续分为食品、补剂和清真三个维度。食品先产生绝对营养结论，再单独显示同类比较提示；候选不足不再降低绝对结论。七个关键营养字段明确为能量、脂肪、饱和脂肪、碳水化合物、糖、蛋白质和盐，纤维未标注不会导致待核对。",
    )

    revision_rows = [
        ("补剂规则", "任何标注维生素 A 的补剂，包括复合维生素和鱼肝油，在疗程内直接判不符合约束"),
        ("营养评分", "新增 0 至 100 分计算公式、分项锚点、品类权重和缺失字段处理"),
        ("品类例外", "油脂按饱和脂肪占总脂肪比例比较；酱料按实际或参考份量评分"),
        ("冷启动", "绝对分数决定值得买或可以考虑；候选库只控制是否显示有更好选择"),
        ("OCR 字段", "明确七个关键字段，纤维作为可选字段；保留快速手工录入兜底"),
        ("清真规则", "补充鱼类海鲜、蛋和混合加工食品，并建立存疑成分词典"),
        ("HealthKit", "第 1 周验证 correlation 更新删除；每个营养样本使用独立同步标识"),
        ("数据模型", "新增容器、待核对队列、评分配置和 HealthKit 同步记录"),
        ("版本命名", "计划书使用修订号，产品继续使用 v0.1、v0.2 等版本号"),
        ("进度兜底", "OCR 未达到 90% 时切换为确认优先或手工优先，不阻塞第 5 周"),
    ]
    replace_table(
        doc,
        ["修订项", "v0.2 决定"],
        ["修订项", "修订 3 决定"],
        revision_rows,
        widths=[1.35, 5.45],
        center_cols=[0],
        font_size=9.2,
    )

    replace_paragraph(
        doc,
        "健康数据：读取步数、活动能量、训练和体重；写入热量、蛋白质、碳水化合物、总脂肪、饱和脂肪和纤维。",
        "健康数据：读取步数、活动能量和训练；写入热量、蛋白质、碳水化合物、总脂肪、饱和脂肪和纤维。v0.1 不读取 HealthKit 体重。",
    )

    success_rows = [
        ("开发周期", "前 6 周完成可安装版本；第 7 至第 8 周只修复阻断问题，不增加功能"),
        ("个人试用", "连续 2 周，每周至少有 5 个完整记录日，并完成至少 2 次超市扫描"),
        ("称重记录", "复用菜谱时 30 秒内完成一餐；首次建菜谱控制在 3 分钟内"),
        ("外卖记录", "常用模板在 30 秒内完成；结果必须显示估值范围和数据来源"),
        ("OCR 质量目标", "30 个标签的七个关键字段在人工修正前达到 90%；未达到时自动采用确认或手工兜底"),
        ("手工兜底", "不依赖 OCR，也能在 60 秒内录完七个字段、配料状态和适用的清真信息"),
        ("安全边界", "已经确认的硬约束不得出现错误绿灯；未经核对的结果不得显示正式绿灯"),
        ("数据一致性", "同一餐新增、编辑或删除后，SwiftData 与 HealthKit 不出现重复记录"),
    ]
    replace_table(
        doc,
        ["指标", "验收标准"],
        ["指标", "验收标准"],
        success_rows,
        widths=[1.35, 5.45],
        center_cols=[0],
        font_size=9.25,
    )

    # OCR fields, completeness, and fallback.
    replace_paragraph(
        doc,
        "3. Vision 在设备上识别表格和文字。解析器统一 kcal、kJ、每份、每 100g、每 100ml、盐和钠。",
        "3. Vision 在设备上识别七个关键营养字段：能量、脂肪、饱和脂肪、碳水化合物、糖、蛋白质和盐。解析器统一 kcal、kJ、每份、每 100g、每 100ml、盐和钠。纤维单独保存为可选字段。",
    )
    replace_paragraph(
        doc,
        "5. 没有明确冲突时检查字段完整度。未完成核对的记录只输出待核对临时结论。",
        "5. 没有明确冲突时检查当前品类需要的字段、配料和认证状态。七个关键字段缺失或尚未确认时返回待核对；英国标签未标纤维时不因此阻断正式结论。",
    )
    replace_paragraph(
        doc,
        "6. 完成核对后计算品类评分，并在有真实候选时判断是否存在更好选择。",
        "6. 完成核对后按品类计算绝对营养分，先给出不建议、可以考虑或值得买，再在有真实同类候选时追加有更好选择提示。",
    )
    replace_paragraph(
        doc,
        "7. 保存原始图片、OCR 候选、人工修正、规则版本和最终结论，便于以后复查。",
        "7. OCR 达不到质量目标时，界面切换为逐字段确认或快速手工录入；该兜底流程必须能独立生成正式结论。",
    )
    scan_anchor = find_paragraph(doc, "2.6 货架前的临时结论")
    insert_block_before(
        doc,
        scan_anchor,
        lambda: add_paragraph(doc, "8. 保存原始图片、OCR 候选、人工修正、规则版本和最终结论，便于以后复查。", after=3),
    )

    temporary_rows = [
        ("疑似不符合约束", "OCR 发现维生素 A 补剂、未认证肉类或其他硬约束证据，但尚未人工确认"),
        ("待核对", "七个关键字段、配料、单位、品类或适用的认证状态尚未确认"),
        ("正式结论", "七个关键字段、配料和适用的认证状态已经确认；纤维缺失不阻断"),
    ]
    replace_table(
        doc,
        ["状态", "触发条件"],
        ["状态", "触发条件"],
        temporary_rows,
        widths=[1.65, 5.15],
        center_cols=[0],
        font_size=9.25,
    )

    logic_rows = [
        ("1", "明确违规证据", "已核验的药物或用户硬约束命中后，直接返回不符合约束"),
        ("2", "数据完整度", "没有明确冲突时检查该产品类型的必需字段；不足时返回待核对"),
        ("3", "产品类型分流", "补剂走独立出口；普通食品继续处理清真标记和营养评分"),
        ("4", "清真标记", "按品类规则生成独立状态，并按用户设置决定是否构成硬约束"),
        ("5", "绝对营养结论", "使用版本化品类权重，生成不建议、可以考虑或值得买"),
        ("6", "同类比较提示", "存在真实候选并达到差值阈值时，追加有更好选择；不覆盖绝对结论"),
    ]
    replace_table(
        doc,
        ["顺序", "检查", "处理"],
        ["顺序", "检查", "处理"],
        logic_rows,
        widths=[0.62, 1.55, 4.63],
        center_cols=[0, 1],
        font_size=9.15,
    )

    food_result_rows = [
        ("硬约束", "不符合约束", "命中已启用、已有来源且已核验的硬约束"),
        ("完整度", "待核对", "缺少该品类的必需字段，或扫描结果尚未人工确认"),
        ("绝对结论", "不建议", "营养分低于 45 分；或大份量品类同时触发两个交通灯高值"),
        ("绝对结论", "可以考虑", "营养分为 45 至 69 分，且没有被高值规则覆盖"),
        ("绝对结论", "值得买", "营养分至少 70 分，且没有被高值规则覆盖"),
        ("比较提示", "有更好选择", "真实同类候选高出至少 15 分，并满足主要指标改善条件"),
    ]
    replace_table(
        doc,
        ["结论", "条件", "优先级"],
        ["输出层", "状态", "条件"],
        food_result_rows,
        widths=[1.15, 1.35, 4.3],
        center_cols=[0, 1],
        font_size=9.05,
    )
    replace_paragraph(
        doc,
        "不建议优先于有更好选择。只有本品已经通过不建议判定后，系统才检查候选差值。分数和阈值是第一版工程基线，不是医疗阈值；完成 50 件真实商品测试后再调整。",
        "候选库不参与绝对营养结论。冷启动时，70 分以上仍显示值得买，45 至 69 分仍显示可以考虑；只有比较条件成立时，结果卡才追加有更好选择。分数和阈值是个人工程基线，不是医疗阈值，完成 50 件真实商品测试后再校准。",
    )

    supplement_rows = [
        ("不符合约束", "异维 A 酸疗程有效，且标签确认该补剂提供维生素 A；包括复合维生素、鱼肝油以及标注 vitamin A、retinol、retinyl acetate、retinyl palmitate 或 vitamin A from beta-carotene 的产品"),
        ("需咨询药师", "产品只标 β 胡萝卜素且未标为维生素 A，维生素 A 形式或来源不明，标签不完整，或其他个人补剂规则尚未由药师确认"),
        ("未发现冲突", "药师已经确认规则范围，标签完整，且没有命中任何已启用冲突；此结论不表示医学安全保证"),
    ]
    replace_table(
        doc,
        ["补剂结论", "使用条件"],
        ["补剂结论", "使用条件"],
        supplement_rows,
        widths=[1.5, 5.3],
        center_cols=[0],
        font_size=8.95,
    )
    replace_paragraph(
        doc,
        "药师确认之前，补剂默认显示需咨询药师，不能显示未发现冲突。NHS 和英国药品说明书均明确要求异维 A 酸治疗期间避免维生素 A 补剂，因此这条规则从第 1 周启用，不等待药师确认。β 胡萝卜素、低剂量复合维生素和强化食品继续保留为需要确认的边缘情况。",
        "药师确认之前，未命中明确红灯的补剂默认显示需咨询药师，不能显示未发现冲突。NHS 将维生素 A 补剂和含维生素 A 的补剂都列为不应与异维 A 酸同用，因此复合维生素和鱼肝油不设低剂量例外。普通强化食品仍按食品逻辑处理；只标 β 胡萝卜素且未标为维生素 A 的补剂保留为需咨询药师。",
    )

    halal_rows = [
        ("已认证", "存在用户认可机构的有效认证"),
        ("成分未见冲突但无认证", "植物性食品、乳制品、蛋，或在用户接受范围内的原形鱼类海鲜，完整配料未发现冲突"),
        ("存疑", "混合加工食品含来源不明的动物性、加工助剂或含酒精香精等成分，需要进一步证明"),
        ("不接受", "肉类或含明胶商品没有用户认可的认证，或发现明确冲突成分"),
    ]
    replace_table(
        doc,
        ["清真标记", "v0.1 定义"],
        ["清真标记", "产品目标 v0.1 定义"],
        halal_rows,
        widths=[2.05, 4.75],
        center_cols=[0],
        font_size=9.15,
    )
    replace_paragraph(
        doc,
        "肉类和含明胶商品必须有用户认可的认证；植物性食品和乳制品在配料完整且未见冲突时可以接受。",
        "肉类和含明胶商品必须有用户认可的认证。原形蛋默认按成分未见冲突处理；原形鱼类海鲜按用户配置的物种范围处理，甲壳类和软体动物不由 App 擅自裁决。加工鱼类、蛋制品、植物性食品和乳制品仍需检查完整配料。",
    )
    replace_paragraph(
        doc,
        "遇到存疑成分时，由用户选择严格模式或允许暂存模式。此配置只表达个人接受标准，不代表 App 作教法裁决。",
        "混合加工食品如含肉、明胶或其衍生成分，必须满足相应认证规则；其他商品进入存疑词典检查。首批存疑词包括凝乳酶、乳清、L-半胱氨酸、胭脂虫红或 E120、虫胶或 E904，以及含酒精的香精。",
    )
    halal_anchor = find_paragraph(doc, "3.5 第一版参考线")
    insert_block_before(
        doc,
        halal_anchor,
        lambda: add_bullet(doc, "遇到存疑成分时，由用户选择严格模式或允许暂存模式。该设置只表达个人接受标准，不代表 App 作教法裁决。"),
    )

    baseline_rows = [
        ("固体食品高值", "饱和脂肪大于 5g/100g；糖大于 22.5g/100g；盐大于 1.5g/100g", "英国交通灯参考线"),
        ("饮料高值", "饱和脂肪大于 2.5g/100ml；糖大于 11.25g/100ml；盐大于 0.75g/100ml", "英国交通灯参考线"),
        ("纤维来源", "至少 3g/100g 或 1.5g/100kcal；高纤维为至少 6g/100g 或 3g/100kcal", "英国营养声称门槛"),
        ("高蛋白声称", "至少 20% 能量来自蛋白质", "英国营养声称门槛"),
        ("个人蛋白密度", "至少 10g/100kcal 为强，6 至 10g 为中，低于 6g 为弱", "个人工程基线"),
        ("油脂", "按饱和脂肪占总脂肪比例评分，只与相同用途油脂比较", "品类工程基线"),
        ("酱料调味品", "按实际份量评分；未设置时使用 15g 参考份量", "品类工程基线"),
        ("更好选择", "候选至少高 15 分，并改善主要指标至少 20%，且饱和脂肪或盐不得恶化超过 10%", "比较提示基线"),
    ]
    replace_table(
        doc,
        ["指标", "第一版阈值", "性质"],
        ["指标", "第一版阈值", "性质"],
        baseline_rows,
        widths=[1.35, 4.25, 1.2],
        center_cols=[0, 2],
        font_size=8.65,
    )
    replace_paragraph(
        doc,
        "冷启动时先使用上述参考线和种子商品。没有真实候选时不显示有更好选择，而是显示可以考虑。评分配置在第 1 周写入版本化 JSON，并为每项保存来源、版本和生效日期。",
        "冷启动时使用绝对分数产生值得买、可以考虑或不建议。候选不足只隐藏有更好选择，不改变绝对结论。评分配置在第 1 周写入版本化 JSON，并为每项保存适用品类、来源、版本和生效日期。",
    )

    medical_heading = find_paragraph(doc, "3.6 医疗边界")

    def add_score_section():
        add_heading(doc, "3.6 营养分计算", 2)
        add_paragraph(
            doc,
            "每个适用分项先换算为 0 至 100 分，再按品类权重计算加权平均。最终分数 = round〔Σ 分项分数 × 权重 ÷ Σ 有效权重〕，并限制在 0 至 100。必需字段缺失时不计算；纤维等可选字段未标注时，将其权重按比例分配给该品类其余有效分项。",
        )
        component_rows = [
            ("蛋白质密度", "不高于 2g/100kcal", "至少 10g/100kcal", "两端之间线性插值"),
            ("纤维密度", "0g/100kcal", "至少 3g/100kcal", "可选字段；缺失时重分权重"),
            ("饱和脂肪", "至少 20% 能量来自饱和脂肪", "不高于 5%", "一般食品按能量占比"),
            ("糖", "达到该固体或饮料高值线", "0", "按每 100g 或 100ml 线性插值"),
            ("盐", "达到该固体或饮料高值线", "0", "按每 100g 或 100ml 线性插值"),
            ("能量密度", "固体 500kcal/100g；饮料 100kcal/100ml", "固体 100kcal/100g；饮料 0", "两端之间线性插值"),
            ("油脂饱和脂肪占比", "至少占总脂肪 25%", "不高于 10%", "只用于油脂品类"),
        ]
        add_table(
            doc,
            ["分项", "0 分锚点", "100 分锚点", "说明"],
            component_rows,
            widths=[1.35, 1.85, 1.75, 1.85],
            center_cols=[0],
            font_size=8.25,
        )
        category_rows = [
            ("蛋白质主菜", "蛋白质 35；纤维 10；饱和脂肪 30；盐 15；能量 10"),
            ("主食和谷物", "蛋白质 15；纤维 30；饱和脂肪 20；糖 20；盐 15"),
            ("混合餐食", "蛋白质 25；纤维 20；饱和脂肪 25；盐 20；能量 10"),
            ("零食和甜点", "蛋白质 10；纤维 20；饱和脂肪 25；糖 30；能量 15"),
            ("饮料", "蛋白质 15；饱和脂肪 20；糖 45；能量 20"),
            ("油脂", "油脂饱和脂肪占比 100"),
            ("酱料和调味品", "每份盐 50；每份糖 25；每份饱和脂肪 15；每份能量 10"),
        ]
        add_table(
            doc,
            ["购买品类", "分项权重"],
            category_rows,
            widths=[1.6, 5.2],
            center_cols=[0],
            font_size=8.75,
        )
        add_paragraph(
            doc,
            "交通灯高值继续展示，但只有参考份量至少 100g 的固体食品或至少 150ml 的饮料同时出现两个高值时，才直接覆盖为不建议。油脂、酱料和调味品不使用每 100g 的双高值覆盖规则；它们分别按脂肪构成和实际份量评分。",
        )

    insert_block_before(doc, medical_heading, add_score_section)
    set_paragraph_text(medical_heading, "3.7 医疗边界")

    replace_paragraph(
        doc,
        "每餐使用 HKCorrelation 表示一个 food 对象，并写入 dietaryEnergyConsumed、dietaryProtein、dietaryCarbohydrates、dietaryFatTotal、dietaryFatSaturated 和 dietaryFiber。每个对象保存 SyncIdentifier 和 SyncVersion，以支持编辑和删除。步数与活动能量使用统计查询读取，不手工累加 iPhone 和 Apple Watch 来源。",
        "每餐使用 HKCorrelation 表示一个 food 对象，并写入 dietaryEnergyConsumed、dietaryProtein、dietaryCarbohydrates、dietaryFatTotal、dietaryFatSaturated 和 dietaryFiber。correlation 使用餐 ID 作为同步标识；每个营养样本使用“餐 ID + 营养素类型”作为独立同步标识，并共享递增版本。第 1 周先实测更新和删除：保存一餐、修改版本、删除 correlation 及其营养样本，再查询确认没有残留或重复。未通过实验前，不假设删除 correlation 会自动删除其中样本。步数和活动能量使用统计查询读取；v0.1 不读取体重。",
    )

    model_rows = [
        ("GoalProfile", "热量、蛋白质、饱和脂肪和纤维目标；来源、生效日期和版本"),
        ("ConstraintRule", "规则类型、证据来源、版本、有效期、严重度和用户确认"),
        ("ScoreProfile", "购买品类、分项权重、锚点、参考份量、版本和生效日期"),
        ("FoodItem", "来源 ID、品类、measurementBasis、参考份量、营养、密度、版本和核对状态"),
        ("ProductSnapshot", "条码、标签原文、配料、营养、价格、店铺、日期和照片"),
        ("ContainerProfile", "容器名称、皮重、单位、最近使用时间和校准日期"),
        ("ReviewQueueItem", "待核对类型、缺失字段、OCR 候选、优先级、状态和来源对象"),
        ("RecipeBatch", "食材、容器皮重、成品重、剩余重、成品率和稳定期"),
        ("TakeawayTemplate", "CoFID 或个人来源、每 100g 营养、范围和餐厅修正"),
        ("MealLog", "餐次、重量、营养快照、估算范围、完整度和 HealthKit 同步版本"),
        ("HealthKitSyncRecord", "餐 ID、对象类型、同步标识、版本、HealthKit UUID 和删除状态"),
        ("DailySnapshot", "目标版本、每日摄入、完整度、滚动统计和 HealthKit 活动"),
        ("Assessment", "绝对食品结论、比较提示、补剂结论、清真标记、分数和证据"),
    ]
    replace_table(
        doc,
        ["模型", "关键内容"],
        ["模型", "关键内容"],
        model_rows,
        widths=[1.65, 5.15],
        center_cols=[0],
        font_size=8.7,
    )

    schedule_rows = [
        ("第 1 周", "基础和规则", "建立项目和 VersionedSchema；写补剂规则、评分公式和品类权重；完成 HealthKit 更新删除实验", "规则测试可运行；HealthKit 实验记录明确；确定付费账号"),
        ("第 2 周", "称重和外卖", "导入 CoFID 子集；实现 FoodItem、ContainerProfile、菜谱批次、个人份量和外卖模板", "可记录真实家常菜和外卖；开始每天使用"),
        ("第 3 周", "今日页", "实现目标差额、完整度、3 天和 7 天完整日平均、稳定期和剩菜复用", "连续 5 天记录真实饮食；修正操作成本"),
        ("第 4 周", "标签扫描", "接入条码、七字段 OCR、快速手工录入、单位换算、人工确认和 ReviewQueueItem", "完成 30 标签基线；无论准确率多少都能手工完成记录"),
        ("第 5 周", "购买判断", "实现违规证据优先、补剂出口、清真扩展、品类评分、绝对结论和比较提示", "无错误绿灯；候选不足不改变绝对结论"),
        ("第 6 周", "HealthKit 和交付", "按实验结论实现读写与更新删除；完成隐私、缓存、导出和 50 件商品测试", "无重复数据；隐私测试通过；进入冻结版本"),
        ("第 7 周", "个人试用一", "真实使用，不增加功能；记录缺失、耗时、错误判断和崩溃", "至少 5 个完整日和 2 次扫描；只修阻断问题"),
        ("第 8 周", "个人试用二", "重复同一验收标准；完成数据导出和恢复演练；整理产品 v0.2 清单", "再次达到 5 个完整日；形成是否继续使用的结论"),
    ]
    replace_table(
        doc,
        ["周期", "目标", "主要任务", "完成定义"],
        ["周期", "目标", "主要任务", "完成定义"],
        schedule_rows,
        widths=[0.72, 1.2, 3.15, 1.73],
        center_cols=[0],
        font_size=8.05,
    )

    test_rows = [
        ("规则顺序", "配料已出现 retinyl palmitate，但营养表缺失", "先返回不符合约束，不被完整度覆盖"),
        ("补剂", "维生素 A 复合维生素、含维生素 A 鱼肝油、仅 β 胡萝卜素产品", "前两类红灯；仅 β 胡萝卜素且未标维 A 时需咨询药师"),
        ("营养分", "每个品类至少 5 个固定样例和边界值", "分项、权重、缺失纤维重分和 45、70 分边界可复算"),
        ("品类例外", "菜籽油、橄榄油、黄油、酱油和低盐酱油", "油脂不因每 100g 高值误伤；酱料按实际份量区分"),
        ("清真", "肉类、鱼类海鲜、蛋、乳制品和含存疑词的混合加工品", "分类和用户配置正确；存疑词不被自动判为不接受"),
        ("OCR", "英国四大超市和亚洲超市的 30 个标签", "记录七字段准确率；低于 90% 时切换确认或手工优先"),
        ("OCR 兜底", "关闭 OCR 或全部识别错误", "60 秒内手工完成七字段和必要配料状态，不阻塞判断"),
        ("称重", "容器皮重、一锅多餐、差值法、比例法、剩菜和复热失水", "个人营养分配守恒；修改后历史正确"),
        ("外卖", "CoFID 模板、整盒重量、容器修正和个人餐厅修正", "显示范围；来源和修正可追溯"),
        ("滚动统计", "完整日、部分完整日和缺失日混合", "不把缺失日当零；样本不足时明确提示"),
        ("HealthKit", "保存、升版本、删除 correlation 和六类营养样本", "每类样本数量正确；查询无残留或重复"),
        ("隐私", "相册照片包含 GPS、EXIF 和文件名", "上传副本不含位置和无关元数据；日志不含 API Key"),
    ]
    replace_table(
        doc,
        ["测试域", "测试集", "通过条件"],
        ["测试域", "测试集", "通过条件"],
        test_rows,
        widths=[1.05, 3.0, 2.75],
        center_cols=[0],
        font_size=8.3,
    )

    risk_rows = [
        ("高精度记录负担过大", "高", "菜谱复用、容器皮重和 3 次稳定期；用真实完成率决定是否继续"),
        ("外卖估值偏差", "高", "使用范围而非单点值；保存同店同菜修正；始终显示来源和置信度"),
        ("众包商品数据错误", "高", "标签优先；保存来源和核对状态；未核对结果不显示绿灯"),
        ("医疗规则误判", "高", "含维生素 A 的补剂明确红灯；其他补剂默认需咨询药师；规则按疗程到期"),
        ("品类评分误伤", "高", "油脂与酱料使用独立评分；每个品类设置固定回归样例"),
        ("清真状态过度推断", "高", "按品类和个人配置处理；存疑词不自动判不接受；不声称代替教法判断"),
        ("比较库为空", "中", "绝对结论不依赖候选；无候选时只隐藏比较提示"),
        ("OCR 未达到 90%", "中", "保留逐字段确认和 60 秒手工录入；不阻塞后续开发"),
        ("HealthKit 删除残留", "高", "第 1 周先做实验；每个营养样本单独同步和追踪；第 6 周按结果实现"),
        ("Open Food Facts 限流", "中", "自定义 User Agent、本地缓存、节流和失败回退到手工录入"),
        ("免费签名失效", "高", "第 1 周确定 Apple Developer Program 账号；日常版本采用付费会员签名"),
        ("本地数据丢失", "中", "第 6 周实现全量 JSON 导出和恢复演练；CloudKit 留到后续决策"),
    ]
    replace_table(
        doc,
        ["风险", "等级", "控制措施"],
        ["风险", "等级", "控制措施"],
        risk_rows,
        widths=[1.85, 0.62, 4.33],
        center_cols=[1],
        font_size=8.55,
    )

    roadmap = find_table(doc, ["版本", "主题", "内容"])
    for row in roadmap.rows[1:]:
        if row.cells[0].text.strip() == "v0.2":
            row.cells[2].text = "下一餐建议、AI 外卖识别、体重趋势、连锁菜单热量、实验室指标时间线和 CloudKit"
            from build_project_plan import format_cell

            format_cell(row.cells[2], center=False, font_size=9.25)

    replace_paragraph(
        doc,
        "定义 GoalProfile、ConstraintRule、FoodItem、ProductSnapshot、MealLog 和 Assessment 第一版模型。",
        "定义 GoalProfile、ConstraintRule、ScoreProfile、FoodItem、ContainerProfile、ReviewQueueItem、MealLog、HealthKitSyncRecord 和 Assessment 第一版模型。",
    )
    replace_paragraph(
        doc,
        "启用异维 A 酸疗程中的维生素 A 补剂规则，并为 β 胡萝卜素等边缘情况保留需咨询药师。",
        "启用异维 A 酸疗程中的维生素 A 补剂规则；复合维生素和含维生素 A 的鱼肝油直接红灯，只标 β 胡萝卜素且未标维生素 A 时需咨询药师。",
    )
    replace_paragraph(
        doc,
        "把英国交通灯、纤维声称、蛋白质声称和个人蛋白密度阈值写入版本化配置。",
        "把营养分公式、分项锚点、七类品类权重、参考份量、英国交通灯和比较阈值写入版本化配置。",
    )
    replace_paragraph(
        doc,
        "建立首批单元测试：明确冲突优先、信息不足、补剂默认状态和清真四态。",
        "建立首批单元测试：明确冲突优先、维生素 A 补剂、评分边界、油脂和酱料例外、信息不足及清真品类。",
    )
    seed_para = find_paragraph(doc, "导入 20 个常用中式食材、4 个 CoFID 外卖模板和 10 个英国包装商品作为种子数据。")
    insert_block_before(
        doc,
        seed_para,
        lambda: (
            add_bullet(doc, "完成 HealthKit 最小实验：保存、更新和删除一餐及六类营养样本，记录查询结果和权限边界。"),
            add_bullet(doc, "完成七字段快速手工录入原型，确保 OCR 关闭时仍能生成正式判断。"),
        ),
    )
    replace_paragraph(
        doc,
        "认可的清真认证机构，以及对存疑成分采用严格模式还是暂存模式。",
        "认可的清真认证机构、可接受的鱼类海鲜范围，以及对甲壳类、软体动物和存疑成分采用严格模式还是暂存模式。",
    )

    disclaimer = find_paragraph(
        doc,
        "本计划中的医学边界用于限制软件行为，不替代医生或药师意见。任何新增红灯规则都必须具备明确来源、适用范围、版本、有效期和用户确认。",
    )

    def add_reference():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.left_indent = Inches(0.12)
        run = paragraph.add_run("14. ")
        set_run_font(run, FONT_BODY, 9.5)
        add_hyperlink(
            paragraph,
            "Apple HealthKit 删除对象",
            "https://developer.apple.com/documentation/healthkit/hkhealthstore/delete(_:withcompletion:)-17hzm",
        )

    insert_block_before(doc, disclaimer, add_reference)

    properties = doc.core_properties
    properties.title = "个人食品决策助手第一阶段项目落实计划书"
    properties.subject = "计划书修订 3 产品目标 v0.1"
    properties.author = ""
    properties.keywords = "iOS SwiftUI 食品标签 中餐称重 外卖估算 HealthKit 修订3"

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    print(build_document().resolve())
