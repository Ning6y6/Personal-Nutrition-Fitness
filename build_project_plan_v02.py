from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_project_plan import (
    COLOR_BLUE,
    COLOR_BORDER,
    COLOR_NAVY,
    COLOR_TEXT_GRAY,
    FONT_BODY,
    FONT_HEAD,
    add_bullet,
    add_heading,
    add_hyperlink,
    add_paragraph,
    add_table,
    format_cell,
    set_cell_shading,
    set_cell_width,
    set_repeat_table_header,
    set_row_cant_split,
    set_run_font,
    set_table_borders,
    setup_styles,
)


OUTPUT_DIR = Path("deliverables")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT_PATH = OUTPUT_DIR / "个人食品决策助手第一阶段项目落实计划书_v0.2.docx"


def add_page_number(paragraph):
    """Add a PAGE field using fldSimple for Word/LibreOffice compatibility."""
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = paragraph.add_run("第 ")
    set_run_font(prefix, FONT_BODY, 8.5, False, COLOR_TEXT_GRAY)

    result_run = paragraph.add_run("1")
    set_run_font(result_run, FONT_BODY, 8.5, False, COLOR_TEXT_GRAY)
    paragraph._p.remove(result_run._r)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    field.append(result_run._r)
    paragraph._p.append(field)

    suffix = paragraph.add_run(" 页")
    set_run_font(suffix, FONT_BODY, 8.5, False, COLOR_TEXT_GRAY)


def add_steps(doc, steps):
    """Add an explicitly numbered list that restarts at 1 for each flow."""
    for index, step in enumerate(steps, 1):
        add_paragraph(doc, f"{index}. {step}", after=3)


def add_metadata_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_borders(table, color=COLOR_BORDER, size="4")

    header = table.add_row()
    set_repeat_table_header(header)
    set_row_cant_split(header)
    for index, text in enumerate(("项目概况", "内容")):
        cell = header.cells[index]
        cell.text = text
        set_cell_width(cell, (1.25, 5.55)[index])
        set_cell_shading(cell, COLOR_NAVY)
        format_cell(cell, header=True, center=True, font_size=9.2)

    for key, value in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        cells[0].text = key
        cells[1].text = value
        set_cell_width(cells[0], 1.25)
        set_cell_width(cells[1], 5.55)
        set_cell_shading(cells[0], COLOR_BLUE)
        format_cell(cells[0], center=True, font_size=9.5)
        format_cell(cells[1], center=False, font_size=9.5)
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    doc.add_paragraph()


def build_document():
    doc = Document()
    setup_styles(doc)

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_before = Pt(28)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("个人食品决策助手第一阶段项目落实计划书")
    set_run_font(run, FONT_HEAD, 25, True, "000000")

    subtitle = doc.add_paragraph(style="Document Subtitle")
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("版本 0.2  原生 iOS 个人自用项目")
    set_run_font(run, FONT_BODY, 12, False, COLOR_TEXT_GRAY)

    add_metadata_table(
        doc,
        [
            ("修订日期", "2026 年 9 月 26 日"),
            ("目标设备", "iPhone 16 Pro Max、iOS 26.6.2、Apple Watch"),
            ("用户目标", "健身小白阶段的减脂增肌，同时处理 LDL 约束、异维 A 酸疗程和清真饮食"),
            ("第一阶段周期", "6 周开发，加 2 周个人试用"),
        ],
    )

    add_heading(doc, "项目结论", 1)
    add_paragraph(
        doc,
        "第一阶段先完成每天都会使用的记录闭环，再完成超市购买判断。第 2 周交付家常菜称重和手工外卖估算，第 3 周交付目标与今日页，从这时开始积累真实数据。第 4 至第 6 周完成标签扫描、约束判断、HealthKit 和导出。开发结束后单独安排 2 周试用，不再把试用压进第 6 周。",
    )
    add_paragraph(
        doc,
        "购买判断分成食品、补剂和清真三个维度。系统先查找明确的违规证据，再检查信息完整度；补剂使用独立结论，清真状态显示为独立标记。未完成人工核对的扫描结果只能显示临时结论，不得显示绿灯。",
    )

    revision_rows = [
        ("周期", "改为 6 周开发加 2 周试用，成功标准放到第 7 至第 8 周验收"),
        ("功能顺序", "称重、外卖估算、目标和今日页提前到第 2 至第 3 周"),
        ("判断顺序", "明确违规证据优先，其次检查完整度，再处理补剂、清真、评分和替代项"),
        ("补剂", "新增独立三档出口，异维 A 酸疗程中的维生素 A 补剂规则从第 1 周启用"),
        ("清真", "独立标记，并按肉类、明胶、植物性食品和乳制品设置不同接受条件"),
        ("冷启动", "加入内置参考线、种子商品和数值化比较阈值"),
        ("完整记录", "增加目标设置、日完整度和滚动平均的缺失日规则"),
        ("技术补充", "加入 EXIF 清理、液体单位、Open Food Facts 限流、开发者账号和 HealthKit 写入范围"),
    ]
    add_table(doc, ["修订项", "v0.2 决定"], revision_rows, widths=[1.35, 5.45], center_cols=[0], font_size=9.4)

    add_heading(doc, "1 项目范围和成功标准", 1)
    add_heading(doc, "1.1 第一阶段用户", 2)
    add_paragraph(
        doc,
        "该版本只服务一个明确用户。用户在英国生活，主要吃中餐家常菜和外卖，愿意称重换取精度，正在进行减脂增肌，并需要把 LDL 偏高、异维 A 酸疗程和清真偏好带入购买与记录流程。用户尚未提出食物过敏，因此 v0.1 不启用过敏原规则；如果后续确认有过敏，再作为硬约束加入。",
    )

    add_heading(doc, "1.2 第一阶段必须完成的闭环", 2)
    for item in [
        "目标设置：保存每日热量目标、蛋白质下限、饱和脂肪上限和纤维下限，并记录来源、确认日期和有效期。",
        "家常菜称重：记录食材重量、容器皮重、成品重量、个人食用重量、剩菜和复热后的修正。",
        "外卖估算：选择 CoFID 或个人模板，按可食重量生成热量和营养范围，不依赖 AI。",
        "今日页：显示当前摄入、距离目标的差额、日完整度以及 3 天和 7 天完整日平均。",
        "超市扫描：条码优先，OCR 作为补充；允许未核对临时结论，并在确认后生成正式判断。",
        "健康数据：读取步数、活动能量、训练和体重；写入热量、蛋白质、碳水化合物、总脂肪、饱和脂肪和纤维。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "1.3 可验证的成功标准", 2)
    success_rows = [
        ("开发周期", "前 6 周完成可安装版本；第 7 至第 8 周只修复阻断问题，不增加功能"),
        ("个人试用", "连续 2 周，每周至少有 5 个完整记录日，并完成至少 2 次超市扫描"),
        ("称重记录", "复用菜谱时 30 秒内完成一餐；首次建菜谱控制在 3 分钟内"),
        ("外卖记录", "常用模板在 30 秒内完成；结果必须显示估值范围和数据来源"),
        ("OCR 准确率", "30 个真实标签中，7 个关键营养字段在人工修正前的字段级准确率至少为 90%"),
        ("安全边界", "已经确认的硬约束不得出现错误绿灯；未经核对的结果不得显示值得买"),
        ("数据一致性", "同一餐新增、编辑或删除后，SwiftData 与 HealthKit 不出现重复记录"),
    ]
    add_table(doc, ["指标", "验收标准"], success_rows, widths=[1.35, 5.45], center_cols=[0], font_size=9.35)

    add_heading(doc, "1.4 第一阶段不做的内容", 2)
    for item in [
        "训练动作库、渐进超负荷和独立 Apple Watch 训练界面。",
        "AI 外卖识别、自动下一餐建议和自动训练日热量调整。下一餐建议移到 v0.2。",
        "社区、好友、公开商品评价和众包审核。",
        "自动解释 LDL、甘油三酯、肝功能或肌肉症状。",
        "在没有已声明过敏的情况下建立过敏原规则。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2 记录流程", 1)
    add_heading(doc, "2.1 目标设置", 2)
    add_steps(
        doc,
        [
            "输入年龄、身高、当前体重、目标体重和每周训练次数。",
            "手工输入或确认每日热量、蛋白质、饱和脂肪和纤维目标。App 不擅自生成医疗阈值。",
            "为每个目标保存来源，例如用户决定、医生要求或营养师建议，并保存确认日期。",
            "目标变更时创建新版本，不覆盖历史记录。当天差额按当日生效的目标计算。",
        ],
    )

    add_heading(doc, "2.2 家常菜称重", 2)
    add_steps(
        doc,
        [
            "选择已有菜谱或新建菜谱，通过大号数字键盘录入食材克数。",
            "选择已保存的容器或锅具，系统自动扣除皮重。",
            "记录成品总重量。成品重量自然包含烹饪失水的影响。",
            "按个人食用重量、整盘前后差值或比例分配个人份量。",
            "如有大量汤汁或油留在盘中，可选择扣除剩余汤汁重量。",
            "保存剩余重量。下次食用时重新称量，并按剩余营养重新分配。",
            "同一菜谱完成 3 次后进入稳定期，保存平均成品率和常用份量，减少重复称重。",
        ],
    )
    formula = add_paragraph(doc, "个人摄入营养 = 全锅营养 × 个人食用重量 ÷ 成品总重量", before=3, after=8)
    formula.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in formula.runs:
        set_run_font(run, "Arial", 11, True, COLOR_NAVY)

    add_heading(doc, "2.3 外卖估算餐", 2)
    add_steps(
        doc,
        [
            "选择 CoFID 种子模板或个人常点外卖模板。首批模板包括鸡肉咖喱、鸡肉炒面、鸡肉抓饭和虾咖喱等。",
            "称量可食部分，或输入包装标示的净重。容器重量未知时可先保存整盒重量，并标记需要修正。",
            "按模板每 100g 营养计算中心值。来源有样本范围时使用原范围；只有单一值时，第一版使用可配置的正负 25% 不确定区间。",
            "用户可按餐厅和菜名保存修正系数。再次点同一餐时优先复用上次确认结果。",
            "外卖日志必须保留模板来源、重量、估值上下限和置信度，不伪装成精确值。",
        ],
    )

    add_heading(doc, "2.4 今日页和记录完整度", 2)
    completeness_rows = [
        ("完整", "当天所有正餐、含热量饮料和零食均已记录；估算餐已确认重量或份量"),
        ("部分完整", "存在估算范围较宽、容器重量未知或用户明确标记的遗漏"),
        ("未完成", "缺少一餐或用户未完成当日确认"),
    ]
    add_table(doc, ["状态", "定义"], completeness_rows, widths=[1.25, 5.55], center_cols=[0], font_size=9.45)
    for item in [
        "今日页显示热量、蛋白质、饱和脂肪和纤维的当前值、目标值与差额。超过上限时显示超出量，不显示负数缺口。",
        "3 天和 7 天视图同时显示日历窗口和完整日平均。未完成日不进入完整日平均，也不按零摄入处理。",
        "3 天窗口少于 2 个完整日、7 天窗口少于 4 个完整日时显示样本不足。",
        "步数、活动能量和 Apple Watch 训练只作背景信息，不直接增加可吃额度。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "2.5 超市标签扫描", 2)
    add_steps(
        doc,
        [
            "扫描条码，并依次查询本地商品库和 Open Food Facts。",
            "未命中或数据不完整时，拍摄营养表、配料表、清真认证标志和可选价签。",
            "Vision 在设备上识别表格和文字。解析器统一 kcal、kJ、每份、每 100g、每 100ml、盐和钠。",
            "先搜索明确违规证据。发现已经核验的冲突时直接返回不符合约束，不因其他字段缺失而降为无法判断。",
            "没有明确冲突时检查字段完整度。未完成核对的记录只输出待核对临时结论。",
            "完成核对后计算品类评分，并在有真实候选时判断是否存在更好选择。",
            "保存原始图片、OCR 候选、人工修正、规则版本和最终结论，便于以后复查。",
        ],
    )

    add_heading(doc, "2.6 货架前的临时结论", 2)
    temporary_rows = [
        ("疑似不符合约束", "OCR 发现维生素 A 补剂、未认证肉类或其他硬约束证据，但尚未人工确认"),
        ("待核对", "没有发现明确冲突，但营养表、配料表、单位或品类尚未确认"),
        ("正式结论", "7 个关键营养字段、配料和适用的认证状态已经确认"),
    ]
    add_table(doc, ["状态", "触发条件"], temporary_rows, widths=[1.65, 5.15], center_cols=[0], font_size=9.35)
    add_paragraph(doc, "临时结论可以提示风险，但不得显示值得买。用户回家后可以进入待核对列表，逐项确认或修正。")

    add_heading(doc, "3 决策引擎", 1)
    add_heading(doc, "3.1 判断顺序", 2)
    logic_rows = [
        ("1", "明确违规证据", "已核验的药物或用户硬约束命中后，直接返回不符合约束"),
        ("2", "数据完整度", "没有明确冲突时，检查当前产品类型要求的字段；不足时返回待核对"),
        ("3", "产品类型分流", "补剂走补剂出口；普通食品继续处理清真标记和营养评分"),
        ("4", "清真标记", "按品类规则生成独立状态，并按用户的严格程度决定是否构成硬约束"),
        ("5", "营养评分", "使用版本化参考线，只与相同购买目的的商品比较"),
        ("6", "替代判断", "存在真实候选且达到差值阈值时，才显示有更好选择"),
    ]
    add_table(doc, ["顺序", "检查", "处理"], logic_rows, widths=[0.62, 1.55, 4.63], center_cols=[0, 1], font_size=9.25)

    add_heading(doc, "3.2 食品结论和优先级", 2)
    food_result_rows = [
        ("不符合约束", "命中已启用、已有来源且已核验的硬约束", "最高"),
        ("待核对", "缺少该品类的必需字段，或扫描结果尚未人工确认", "2"),
        ("不建议", "营养分低于 45 分，或同时触发两个英国交通灯高值", "3"),
        ("有更好选择", "本品不属于不建议，且真实同类候选高出至少 15 分", "4"),
        ("值得买", "营养分至少 70 分、没有高值项，且没有达到差值阈值的更好候选", "5"),
        ("可以考虑", "营养分为 45 至 69 分，或冷启动阶段没有足够候选可比", "6"),
    ]
    add_table(doc, ["结论", "条件", "优先级"], food_result_rows, widths=[1.4, 4.75, 0.65], center_cols=[0, 2], font_size=9.1)
    add_paragraph(
        doc,
        "不建议优先于有更好选择。只有本品已经通过不建议判定后，系统才检查候选差值。分数和阈值是第一版工程基线，不是医疗阈值；完成 50 件真实商品测试后再调整。",
    )

    add_heading(doc, "3.3 补剂的独立出口", 2)
    supplement_rows = [
        ("不符合约束", "异维 A 酸疗程有效，且完整配料或 Supplement Facts 已确认含维生素 A、retinol、retinyl acetate 或 retinyl palmitate 等已启用成分"),
        ("需咨询药师", "药师尚未确认个人补剂规则，或产品含 β 胡萝卜素、低剂量复合维生素、鱼肝油、来源不明的维生素 A 形式或标签不完整"),
        ("未发现冲突", "药师已经确认规则范围，标签完整，且没有命中任何已启用冲突；此结论不表示医学安全保证"),
    ]
    add_table(doc, ["补剂结论", "使用条件"], supplement_rows, widths=[1.5, 5.3], center_cols=[0], font_size=9.1)
    add_paragraph(
        doc,
        "药师确认之前，补剂默认显示需咨询药师，不能显示未发现冲突。NHS 和英国药品说明书均明确要求异维 A 酸治疗期间避免维生素 A 补剂，因此这条规则从第 1 周启用，不等待药师确认。β 胡萝卜素、低剂量复合维生素和强化食品继续保留为需要确认的边缘情况。",
    )
    add_paragraph(
        doc,
        "蛋白粉和代餐奶昔同时生成两张结果卡。第一张按食品逻辑评估热量、蛋白质和饱和脂肪；第二张按补剂逻辑检查维生素和活性成分。任何补剂红灯都会覆盖营养评分。",
    )

    add_heading(doc, "3.4 清真状态", 2)
    halal_rows = [
        ("已认证", "存在用户认可机构的有效认证"),
        ("成分未见冲突但无认证", "植物性食品或乳制品的完整配料没有发现冲突"),
        ("存疑", "香精、乳化剂、加工助剂或来源不清的成分需要进一步证明"),
        ("不接受", "肉类或含明胶商品没有用户认可的认证，或发现明确冲突成分"),
    ]
    add_table(doc, ["清真标记", "v0.1 定义"], halal_rows, widths=[2.05, 4.75], center_cols=[0], font_size=9.35)
    for item in [
        "清真标记显示在食品结论旁边，不再把认证来源不明自动映射为无法判断。",
        "肉类和含明胶商品必须有用户认可的认证；植物性食品和乳制品在配料完整且未见冲突时可以接受。",
        "遇到存疑成分时，由用户选择严格模式或允许暂存模式。此配置只表达个人接受标准，不代表 App 作教法裁决。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "3.5 第一版参考线", 2)
    baseline_rows = [
        ("固体食品高值", "饱和脂肪大于 5g/100g；糖大于 22.5g/100g；盐大于 1.5g/100g", "英国交通灯参考线"),
        ("饮料高值", "饱和脂肪大于 2.5g/100ml；糖大于 11.25g/100ml；盐大于 0.75g/100ml", "英国交通灯参考线"),
        ("纤维来源", "至少 3g/100g 或 1.5g/100kcal；高纤维为至少 6g/100g 或 3g/100kcal", "英国营养声称门槛"),
        ("高蛋白声称", "至少 20% 能量来自蛋白质", "英国营养声称门槛"),
        ("个人蛋白密度", "至少 10g/100kcal 为强，6 至 10g 为中，低于 6g 为弱", "个人工程基线"),
        ("更好选择", "同品类候选至少高 15 分，并改善一项主要指标至少 20%，且饱和脂肪或盐不得恶化超过 10%", "个人工程基线"),
    ]
    add_table(doc, ["指标", "第一版阈值", "性质"], baseline_rows, widths=[1.35, 4.25, 1.2], center_cols=[0, 2], font_size=8.8)
    add_paragraph(
        doc,
        "冷启动时先使用上述参考线和种子商品。没有真实候选时不显示有更好选择，而是显示可以考虑。评分配置在第 1 周写入版本化 JSON，并为每项保存来源、版本和生效日期。",
    )

    add_heading(doc, "3.6 医疗边界", 2)
    add_paragraph(
        doc,
        "普通食物天然含有维生素 A 时不自动禁食。NHS 说明异维 A 酸治疗期间通常可以正常饮食，但应避免维生素 A 补剂。疗程规则必须保存开始日期和结束日期；疗程结束后自动失效，并提示用户重新确认。实验室数据只用于记录和与医生沟通，不进入自动诊断。",
    )

    add_heading(doc, "4 数据和技术架构", 1)
    add_heading(doc, "4.1 技术栈", 2)
    tech_rows = [
        ("界面", "SwiftUI、Observation、Charts"),
        ("业务数据", "SwiftData、VersionedSchema、本地优先"),
        ("扫描", "Vision RecognizeDocumentsRequest、条码检测、PhotosUI"),
        ("健康数据", "HealthKit、统计查询、food correlation、同步标识"),
        ("网络", "URLSession、Open Food Facts、USDA FoodData Central"),
        ("AI", "FoodVisionProvider 协议，仅为困难图片和后续 AI 外卖识别预留"),
        ("安全", "Keychain 保存 API Key，图片重编码后再上传"),
        ("导出", "JSON 全量备份，CSV 饮食和商品记录"),
    ]
    add_table(doc, ["层", "选择"], tech_rows, widths=[1.25, 5.55], center_cols=[0], font_size=9.35)

    add_heading(doc, "4.2 数据源", 2)
    source_rows = [
        ("CoFID", "英国常见食材、成品菜和外卖模板", "筛选后本地导入；保存数据版本和原始食物代码"),
        ("USDA FoodData Central", "通用食材和补充微量营养", "远程搜索后保存快照；API Key 放入 Keychain"),
        ("Open Food Facts", "英国条码商品", "条码查询优先；保存原始来源、响应时间和人工核对状态"),
        ("个人食材和菜谱", "中式食材、外卖和常做菜", "最高优先级；允许纠错、版本化和复用"),
        ("商品标签", "具体品牌的配料和营养", "产品特定判断时优先于外部数据库"),
    ]
    add_table(doc, ["来源", "用途", "落地要求"], source_rows, widths=[1.45, 2.0, 3.35], center_cols=[0], font_size=9.0)
    add_paragraph(
        doc,
        "已核对 CoFID 2021 数据集，其中包含 Chicken curry average takeaway、Chicken chow mein takeaway、Chicken biryani takeaway 和 Prawn curry takeaway 等条目，可用于外卖模板的首批种子数据。",
    )

    add_heading(doc, "4.3 网络和隐私要求", 2)
    privacy_rows = [
        ("图片上传", "使用 ImageIO 重新编码为新 JPEG 或 HEIC，只保留方向和必要像素；删除 EXIF、GPS、IPTC 和原始文件名"),
        ("Open Food Facts", "设置 AppName/Version ContactEmail 格式的自定义 User Agent；商品查询不超过每分钟 15 次，搜索不超过每分钟 10 次"),
        ("缓存", "条码结果和搜索结果写入本地缓存；命中缓存时不重复请求，30 天后提示刷新而不是强制刷新"),
        ("第三方 AI", "只上传完成当前识别任务所需的裁剪图片，不发送 LDL、用药、清真设置、体重或完整用户档案"),
        ("日志", "不记录 API Key、原始医疗信息、完整配料照片路径或照片位置元数据"),
    ]
    add_table(doc, ["项目", "要求"], privacy_rows, widths=[1.35, 5.45], center_cols=[0], font_size=9.05)

    add_heading(doc, "4.4 重量和体积单位", 2)
    for item in [
        "FoodItem 保存 measurementBasis，取值为 per100g、per100ml 或 perServing。",
        "牛奶和饮料优先直接记录毫升，不把克数静默当作毫升。",
        "只有在产品或食材存在 DensityProfile 时才执行克与毫升换算，并保存密度来源和温度说明。",
        "用户使用厨房秤称液体时，界面要求选择已知密度或改为量杯输入毫升。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "4.5 HealthKit 范围", 2)
    add_paragraph(
        doc,
        "每餐使用 HKCorrelation 表示一个 food 对象，并写入 dietaryEnergyConsumed、dietaryProtein、dietaryCarbohydrates、dietaryFatTotal、dietaryFatSaturated 和 dietaryFiber。每个对象保存 SyncIdentifier 和 SyncVersion，以支持编辑和删除。步数与活动能量使用统计查询读取，不手工累加 iPhone 和 Apple Watch 来源。",
    )

    add_heading(doc, "4.6 核心数据模型", 2)
    model_rows = [
        ("GoalProfile", "热量、蛋白质、饱和脂肪和纤维目标；来源、生效日期和版本"),
        ("ConstraintRule", "规则类型、证据来源、版本、有效期、严重度和用户确认"),
        ("FoodItem", "来源 ID、品类、measurementBasis、营养、密度、版本和核对状态"),
        ("ProductSnapshot", "条码、标签原文、配料、营养、价格、店铺、日期和照片"),
        ("RecipeBatch", "食材、容器皮重、成品重、剩余重、成品率和稳定期"),
        ("TakeawayTemplate", "CoFID 或个人来源、每 100g 营养、范围和餐厅修正"),
        ("MealLog", "餐次、重量、营养快照、估算范围、完整度和 HealthKit 同步版本"),
        ("DailySnapshot", "目标版本、每日摄入、完整度、滚动统计和 HealthKit 活动"),
        ("Assessment", "食品结论、补剂结论、清真标记、命中规则、分数和证据"),
    ]
    add_table(doc, ["模型", "关键内容"], model_rows, widths=[1.65, 5.15], center_cols=[0], font_size=8.95)

    add_heading(doc, "5 八周实施计划", 1)
    schedule_rows = [
        ("第 1 周", "基础和规则", "建立 Xcode 项目、VersionedSchema、目标设置、规则 JSON、阈值表和开发者账号；启用维生素 A 补剂规则", "可设置目标；规则测试可运行；确定付费开发者账号"),
        ("第 2 周", "称重和外卖", "导入 CoFID 子集；实现 FoodItem、菜谱批次、容器皮重、成品重、个人份量和外卖模板", "可记录真实家常菜和外卖；开始每天使用"),
        ("第 3 周", "今日页", "实现目标差额、完整度、3 天和 7 天完整日平均、稳定期和剩菜复用", "连续 5 天记录真实饮食；修正操作成本"),
        ("第 4 周", "标签扫描", "接入条码、Vision OCR、7 个关键字段解析、单位换算、人工确认和待核对队列", "30 个标签的字段级准确率至少 90%"),
        ("第 5 周", "购买判断", "实现违规证据优先、补剂出口、清真标记、评分、种子候选和并排比较", "无错误绿灯；冷启动时不虚构更好选择"),
        ("第 6 周", "HealthKit 和交付", "实现读写、同步更新删除、EXIF 清理、Open Food Facts 缓存、JSON 和 CSV 导出；完成 50 件商品测试", "无重复数据；隐私测试通过；进入冻结版本"),
        ("第 7 周", "个人试用一", "真实使用，不增加功能；记录缺失、耗时、错误判断和崩溃", "至少 5 个完整日和 2 次扫描；只修阻断问题"),
        ("第 8 周", "个人试用二", "重复同一验收标准；完成数据导出和恢复演练；整理 v0.2 清单", "再次达到 5 个完整日；形成是否继续使用的结论"),
    ]
    add_table(
        doc,
        ["周期", "目标", "主要任务", "完成定义"],
        schedule_rows,
        widths=[0.72, 1.2, 3.15, 1.73],
        center_cols=[0],
        font_size=8.25,
    )

    add_heading(doc, "5.1 每周交付纪律", 2)
    for item in [
        "第 1 周使用固定样例验证模型和规则；从第 2 周开始，每周末必须用真实做饭和外卖数据试用。",
        "每周只交付一个可以从界面走通的纵向闭环，不只完成孤立的数据模型。",
        "规则、单位换算和称重公式必须有单元测试；医疗规则还要记录依据和有效期。",
        "第 7 至第 8 周冻结功能，只修复会阻断真实使用或造成错误判断的问题。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "6 验收和测试", 1)
    test_rows = [
        ("规则顺序", "配料已出现 retinyl palmitate，但营养表缺失", "先返回不符合约束，不被完整度覆盖"),
        ("补剂", "维生素 A 补剂、β 胡萝卜素、蛋白粉和代餐奶昔", "三档补剂出口正确；未获药师确认时不显示未发现冲突"),
        ("清真", "有认证肉类、无认证肉类、含明胶商品、植物性食品和乳制品", "独立标记正确；认证不明不再统一变成无法判断"),
        ("OCR", "英国四大超市和亚洲超市的 30 个标签，包含反光和弯曲包装", "7 个关键字段在人工修正前的字段级准确率至少 90%"),
        ("称重", "容器皮重、一锅多餐、差值法、比例法、剩菜和复热失水", "个人营养分配守恒；修改后历史正确"),
        ("外卖", "CoFID 模板、整盒重量、容器修正和个人餐厅修正", "显示范围；来源和修正可追溯"),
        ("滚动统计", "完整日、部分完整日和缺失日混合", "不把缺失日当零；样本不足时明确提示"),
        ("HealthKit", "同一餐新增、编辑和删除；Watch 与 iPhone 同时有活动数据", "营养样本无重复；活动数据不重复相加"),
        ("隐私", "相册照片包含 GPS、EXIF 和文件名", "上传副本不含位置和无关元数据；日志不含 API Key"),
    ]
    add_table(doc, ["测试域", "测试集", "通过条件"], test_rows, widths=[1.05, 3.0, 2.75], center_cols=[0], font_size=8.65)

    add_heading(doc, "6.1 发布前阻断条件", 2)
    for item in [
        "任何已确认的硬约束存在错误绿灯。",
        "未完成人工核对的商品显示值得买。",
        "补剂在药师规则尚未确认时显示未发现冲突。",
        "缺失日被当作零摄入并进入滚动平均。",
        "同一餐编辑后在 HealthKit 出现重复记录。",
        "第三方请求包含用药、LDL、清真设置、体重档案或照片位置元数据。",
        "计算结果不能追溯到食材输入、重量、目标版本和数据库版本。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "7 风险和控制", 1)
    risk_rows = [
        ("高精度记录负担过大", "高", "菜谱复用、容器皮重和 3 次稳定期；第 7 至第 8 周用真实完成率决定是否继续"),
        ("外卖估值偏差", "高", "使用范围而非单点值；保存同店同菜修正；始终显示来源和置信度"),
        ("众包商品数据错误", "高", "标签优先；保存来源和核对状态；未核对结果不显示绿灯"),
        ("医疗规则误判", "高", "只启用有正式来源的规则；补剂默认需咨询药师；规则按疗程到期"),
        ("清真状态过度推断", "高", "独立标记；按品类和个人接受标准处理；不声称代替教法判断"),
        ("比较库为空", "中", "使用阈值和种子商品；无候选时显示可以考虑，不虚构替代项"),
        ("Open Food Facts 限流", "中", "自定义 User Agent、本地缓存、节流和失败回退到手工录入"),
        ("免费签名失效", "高", "第 1 周确定 Apple Developer Program 账号；日常使用版本采用付费会员签名"),
        ("本地数据丢失", "中", "第 6 周实现全量 JSON 导出和恢复演练；CloudKit 留到后续决策"),
    ]
    add_table(doc, ["风险", "等级", "控制措施"], risk_rows, widths=[1.85, 0.62, 4.33], center_cols=[1], font_size=8.8)

    add_heading(doc, "8 后续版本", 1)
    roadmap_rows = [
        ("v0.2", "建议和同步", "下一餐建议、AI 外卖识别、连锁菜单热量、实验室指标时间线和 CloudKit"),
        ("v0.3", "训练联动", "动作模板、重量次数、RPE、AlarmKit 休息计时、Watch 训练重叠检查和训练日饮食调整"),
        ("v0.4", "效率和扩展", "语音称重、App Intents、控制中心入口、更多地区标签和更成熟的替代推荐"),
    ]
    add_table(doc, ["版本", "主题", "内容"], roadmap_rows, widths=[0.72, 1.55, 4.53], center_cols=[0], font_size=9.25)

    add_heading(doc, "9 第一周启动清单", 1)
    first_week = [
        "创建原生 SwiftUI Xcode 项目，确定 Bundle Identifier，并启用 Git。",
        "确认 Apple Developer Program 付费账号。免费 Personal Team 的设备安装签名 7 天到期，不适合每天使用。",
        "建立 Core、Data、Features、Integrations、Resources 和 Tests 六个目录。",
        "定义 GoalProfile、ConstraintRule、FoodItem、ProductSnapshot、MealLog 和 Assessment 第一版模型。",
        "启用 SwiftData VersionedSchema v1，并为未来 CloudKit 保留可迁移结构。",
        "完成目标设置页面的静态原型，支持热量、蛋白质、饱和脂肪和纤维。",
        "完成规则 JSON，写入判断优先级、来源 URL、版本、生效日期、失效日期和用户确认。",
        "启用异维 A 酸疗程中的维生素 A 补剂规则，并为 β 胡萝卜素等边缘情况保留需咨询药师。",
        "把英国交通灯、纤维声称、蛋白质声称和个人蛋白密度阈值写入版本化配置。",
        "建立首批单元测试：明确冲突优先、信息不足、补剂默认状态和清真四态。",
        "导入 20 个常用中式食材、4 个 CoFID 外卖模板和 10 个英国包装商品作为种子数据。",
    ]
    for item in first_week:
        add_bullet(doc, item)

    add_heading(doc, "9.1 启动前需要补齐的个人数据", 2)
    for item in [
        "年龄、身高、当前体重、目标体重和每周训练次数。",
        "每日热量、蛋白质、饱和脂肪和纤维目标；如来自医生或营养师，应保存原文。",
        "异维 A 酸疗程的开始日期、预计结束日期和药品说明书版本。",
        "认可的清真认证机构，以及对存疑成分采用严格模式还是暂存模式。",
        "最常做的 10 道菜、最常点的 10 份外卖和最常购买的 20 件商品。",
        "是否存在任何食物过敏；如没有，v0.1 不启用过敏原检查。",
    ]:
        add_bullet(doc, item)

    add_heading(doc, "依据和技术参考", 1)
    references = [
        ("NHS 异维 A 酸说明", "https://www.nhs.uk/medicines/isotretinoin-roaccutane/"),
        ("英国异维 A 酸药品特性摘要", "https://www.medicines.org.uk/emc/product/100786/smpc"),
        ("英国 CoFID 数据集", "https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid"),
        ("英国前包装营养标签指南", "https://www.gov.uk/government/publications/front-of-pack-nutrition-labelling-guidance"),
        ("英国营养和健康声称法规附件", "https://www.legislation.gov.uk/eur/2006/1924/annexes/2020-12-31"),
        ("Open Food Facts API", "https://openfoodfacts.github.io/documentation/docs/Product-Opener/api/"),
        ("Apple 开发者账号说明", "https://developer.apple.com/help/account/basics/about-your-developer-account"),
        ("Apple iOS 能力支持表", "https://developer.apple.com/help/account/reference/supported-capabilities-ios"),
        ("Apple HealthKit 营养类型", "https://developer.apple.com/documentation/healthkit/nutrition-type-identifiers"),
        ("Apple HealthKit HKCorrelation", "https://developer.apple.com/documentation/healthkit/hkcorrelation"),
        ("Apple HealthKit Sync Identifier", "https://developer.apple.com/documentation/healthkit/hkmetadatakeysyncidentifier"),
        ("Apple Vision RecognizeDocumentsRequest", "https://developer.apple.com/documentation/vision/recognizedocumentsrequest"),
        ("USDA FoodData Central API", "https://fdc.nal.usda.gov/api-guide/"),
    ]
    for index, (name, url) in enumerate(references, 1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.left_indent = Inches(0.12)
        run = paragraph.add_run(f"{index}. ")
        set_run_font(run, FONT_BODY, 9.5)
        add_hyperlink(paragraph, name, url)

    add_paragraph(
        doc,
        "本计划中的医学边界用于限制软件行为，不替代医生或药师意见。任何新增红灯规则都必须具备明确来源、适用范围、版本、有效期和用户确认。",
        before=8,
        after=0,
    )

    properties = doc.core_properties
    properties.title = "个人食品决策助手第一阶段项目落实计划书"
    properties.subject = "原生 iOS 个人食品决策助手 v0.2 实施计划"
    properties.author = ""
    properties.keywords = "iOS SwiftUI 食品标签 中餐称重 外卖估算 HealthKit"

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


if __name__ == "__main__":
    path = build_document()
    print(path.resolve())
